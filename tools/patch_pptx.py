"""
Patch EN PLACE de soutenance/support_eckert.pptx (ne régénère pas le fichier,
préserve la mise en forme de l'auteur) :
  - met à jour les puces des diapos 6 à 10 (retours de revue) ;
  - vide les notes de commentaire de toutes les diapos (le « à dire » vit
    désormais dans conducteur_soutenance.md).

Identification robuste : la diapo est repérée par son titre (pas par son index),
la zone de puces par la présence du caractère « • ».

    PYTHONPATH=. python -m tools.patch_pptx
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

FICHIER = Path(__file__).resolve().parents[1] / "soutenance" / "support_eckert.pptx"

# Nouvelles puces, repérées par un fragment du titre de la diapo.
PUCES = {
    "Trois scénarios de coût": [
        "Local : logiciel gratuit, mais ~45 €/mois de temps humain (+ ~1 350 € de build)",
        "Cloud Azure (ressources détaillées) : ~30 – 90 €/mois",
        "SaaS : Snowflake ~40-80 € (usage) · Fabric ~180-305 € (capacité)",
        "SaaS écarté (réel) : coût récurrent + souveraineté des données de santé",
    ],
    "Un entrepôt en quatre couches": [
        "bronze (brut) · silver (typé) · gold (métier) · ops (journal)",
        "Trois rôles : administration (DBA), écriture (pipeline), lecture (métier)",
        "La lecture ne peut ni écrire ni lire le brut — TESTÉ",
        "Le métier consomme une vue déjà priorisée",
    ],
    "Trois méthodes de pipeline": [
        "Choix ELT : charger le brut, puis transformer (rejouabilité + au plus près)",
        "Fil de l'eau (SQL) : typage, nettoyage, déduplication",
        "Orchestration (Airflow) : enchaîne et fiabilise",
        "Distribué (Spark) : absorbe le volume du rapprochement",
    ],
    "Orchestration : le DAG réel": [
        "4 tâches : catalogue data.gouv → validation contrat → bronze → silver",
        "Retries · SLA · callbacks d'alerte · exécution réelle sur le fichier INSEE",
    ],
    "Distribué : Spark détecte": [
        "Postgres stocke ; Spark lit/écrit via JDBC · INSEE partiel reconstruit",
        "3 niveaux · 4 162 contrats · 88 M€ certain · Spark justifié par le cumul",
    ],
}


def _titre(slide) -> str:
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() and "•" not in sh.text_frame.text:
            return sh.text_frame.text.strip()
    return ""


def _zone_puces(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and "•" in sh.text_frame.text:
            return sh.text_frame
    return None


def _remplacer_puces(tf, lignes) -> None:
    # Récupère la mise en forme de la première puce existante pour la conserver.
    ref = tf.paragraphs[0].runs[0].font if tf.paragraphs and tf.paragraphs[0].runs else None
    taille = ref.size if ref else Pt(18)
    nom = ref.name if ref else "Calibri"
    couleur = None
    if ref is not None:
        try:
            couleur = ref.color.rgb
        except Exception:
            couleur = None
    tf.clear()
    for i, ligne in enumerate(lignes):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = "•  " + ligne
        run.font.size = taille
        run.font.name = nom
        if couleur is not None:
            run.font.color.rgb = couleur


def main() -> None:
    prs = Presentation(str(FICHIER))
    modifiees, notes_videes = [], 0
    for slide in prs.slides:
        titre = _titre(slide)
        for fragment, lignes in PUCES.items():
            if fragment in titre:
                tf = _zone_puces(slide)
                if tf is not None:
                    _remplacer_puces(tf, lignes)
                    modifiees.append(fragment)
                break
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            slide.notes_slide.notes_text_frame.text = ""
            notes_videes += 1
    prs.save(str(FICHIER))
    print("Diapos dont les puces ont ete mises a jour:", len(modifiees))
    print("Notes videes:", notes_videes)


if __name__ == "__main__":
    main()
